"""A4 letterheads taken from each fund's official stationery scan.

Every upload is converted to A4 with a 3.75 cm top margin and a 2.75 cm
bottom margin. Header and footer strips are cropped from the fund's full-page
letterhead PNG (``<slug>_letterhead.png``); Fund1.pdf is only a fallback.
Operators can still drag those strips.
"""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont

from app.config import FUNDS, LETTERHEADS_DIR

# ISO A4 in PDF points, matching Word's "A4" paper size.
A4_WIDTH = 595.28
A4_HEIGHT = 841.89
CM_TO_PT = 72.0 / 2.54
MARGIN_TOP_PT = 3.75 * CM_TO_PT
MARGIN_BOTTOM_PT = 2.75 * CM_TO_PT
MARGIN_SIDE_PT = 2.0 * CM_TO_PT

# Kept for tests and any remaining callers.
WORD_PAPERS: dict[str, tuple[float, float]] = {
    "Letter": (612.0, 792.0),
    "Legal": (612.0, 1008.0),
    "A4": (A4_WIDTH, A4_HEIGHT),
    "A5": (419.53, 595.28),
    "A3": (841.89, 1190.55),
    "Executive": (522.0, 756.0),
    "Tabloid": (792.0, 1224.0),
    "Statement": (396.0, 612.0),
}

PAGE_SETUPS: dict[str, tuple[float, float]] = {}
for _paper, (_w, _h) in WORD_PAPERS.items():
    PAGE_SETUPS[f"{_paper} Portrait"] = (_w, _h)
    PAGE_SETUPS[f"{_paper} Landscape"] = (_h, _w)

DEFAULT_PAGE_SETUP = "A4 Portrait"
MASTER_SCALE = 3.0

_WINDOWS_FONTS = Path(r"C:\Windows\Fonts")
TIMES_REGULAR = _WINDOWS_FONTS / "times.ttf"
TIMES_BOLD = _WINDOWS_FONTS / "timesbd.ttf"

SOURCE_PDF = LETTERHEADS_DIR / "fund_1_source.pdf"
DOWNLOADS_PDF = Path(r"c:\Users\aksha\Downloads\Fund1.pdf")

TRUSTEE_DELHI = [
    "Trustee: Vistra ITCL (India) Limited",
    "Regd. Add.: E-287, 2nd Floor,",
    "Greater Kailash-I, South Delhi,",
    "New Delhi, India - 110048",
]

FUND_HEADERS: dict[str, dict[str, object]] = {
    "Fund 1": {
        "title": "YourNest Angel Fund",
        "lines": [
            "C/o YourNest Capital Advisors P. Limited",
            "Regd. Office : 902, Block 14, Heritage City,",
            "Near MG Road Metro Station,",
            "Gurgaon - 122002",
        ],
    },
    "Fund 2": {
        "title": "YourNest Venture Capital Fund II",
        "lines": TRUSTEE_DELHI,
    },
    "Fund 3": {
        "title": "YourNest Innovative Products VC Fund III",
        "lines": TRUSTEE_DELHI,
    },
    "Fund 4": {
        "title": "YourNest Innovative Products VC Fund IV",
        "lines": TRUSTEE_DELHI,
    },
    "Continuation Fund": {
        "title": "YourNest Continuation Fund",
        "lines": TRUSTEE_DELHI,
    },
    "AMC": {
        "title": "YourNest Capital Advisors Pvt Ltd",
        "lines": [
            "416, 4th Floor, MGF Metropolis,",
            "Near M.G. Road Metro Station,",
            "Gurgaon-122002, Haryana (INDIA)",
        ],
    },
}

LOGO_PATH = LETTERHEADS_DIR / "yournest_logo.png"

# Address block on the 3x Fund 1 scan (leave the red bar untouched).
_ADDR_WHITEN = (1035, 8, 1710, 265)
_ADDR_RIGHT_X = 1651
_ADDR_TOP_Y = 20
_ADDR_LEADING = 37
_ADDR_FONT_PX = 31


def content_box(page_width: float, page_height: float) -> fitz.Rect:
    """Printable body — 3.75 cm from the top, 2.75 cm from the bottom."""
    top = min(MARGIN_TOP_PT, page_height * 0.22)
    bottom = page_height - min(MARGIN_BOTTOM_PT, page_height * 0.16)
    side = min(MARGIN_SIDE_PT, page_width * 0.12)
    if bottom - top < 120:
        top = min(top, page_height * 0.12)
        bottom = page_height - min(MARGIN_BOTTOM_PT, page_height * 0.08)
    return fitz.Rect(side, top, page_width - side, bottom)


def reserved_rects(page_width: float, page_height: float) -> list[fitz.Rect]:
    body = content_box(page_width, page_height)
    return [
        fitz.Rect(0, 0, page_width, body.y0),
        fitz.Rect(0, body.y1, page_width, page_height),
    ]


def letterhead_slug(fund_name: str) -> str:
    return Path(FUNDS[fund_name]).stem


def letterhead_header_path(fund_name: str) -> Path:
    return LETTERHEADS_DIR / f"{letterhead_slug(fund_name)}_header.png"


def letterhead_footer_path(fund_name: str) -> Path:
    return LETTERHEADS_DIR / f"{letterhead_slug(fund_name)}_footer.png"


def letterhead_source_path(fund_name: str) -> Path:
    """Full-page official letterhead scan for the fund."""
    return LETTERHEADS_DIR / f"{letterhead_slug(fund_name)}_letterhead.png"


def ensure_letterhead_assets() -> None:
    LETTERHEADS_DIR.mkdir(parents=True, exist_ok=True)
    _ensure_source_pdf()
    for fund_name in FUNDS:
        header = letterhead_header_path(fund_name)
        footer = letterhead_footer_path(fund_name)
        source = letterhead_source_path(fund_name)
        stale = source.exists() and header.exists() and (
            header.stat().st_mtime < source.stat().st_mtime
        )
        if not header.exists() or not footer.exists() or stale:
            render_letterhead_strips(fund_name)
    if not LOGO_PATH.exists():
        _extract_logo_from_master()


def _ensure_source_pdf() -> Path:
    LETTERHEADS_DIR.mkdir(parents=True, exist_ok=True)
    if SOURCE_PDF.exists() and SOURCE_PDF.stat().st_size > 0:
        return SOURCE_PDF
    if DOWNLOADS_PDF.exists():
        SOURCE_PDF.write_bytes(DOWNLOADS_PDF.read_bytes())
        return SOURCE_PDF
    raise FileNotFoundError(
        "Official Fund 1 letterhead PDF is missing. Place Fund1.pdf at "
        f"{SOURCE_PDF} or {DOWNLOADS_PDF}."
    )


def _load_master_page() -> tuple[Image.Image, float]:
    doc = fitz.open(_ensure_source_pdf())
    try:
        page = doc[0]
        pix = page.get_pixmap(matrix=fitz.Matrix(MASTER_SCALE, MASTER_SCALE), alpha=False)
        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        return image, page.rect.height
    finally:
        doc.close()


def _band_crops(master: Image.Image, pdf_height: float) -> tuple[Image.Image, Image.Image]:
    width, height = master.size
    header_h = max(1, int(round(MARGIN_TOP_PT / pdf_height * height)))
    footer_h = max(1, int(round(MARGIN_BOTTOM_PT / pdf_height * height)))
    header = master.crop((0, 0, width, header_h))
    footer = master.crop((0, height - footer_h, width, height))
    return header, footer


def _crop_full_page_strips(source: Path) -> tuple[Image.Image, Image.Image]:
    """Cut the top/bottom margin bands from a full-page letterhead scan.

    Bands are sized from the image width so the strips keep the A4 aspect
    and are not distorted when painted across the page.
    """
    with Image.open(source) as raw:
        image = raw.convert("RGB")
    width, height = image.size
    px_per_pt = width / A4_WIDTH
    header_h = min(height, max(1, int(round(MARGIN_TOP_PT * px_per_pt))))
    footer_h = min(height, max(1, int(round(MARGIN_BOTTOM_PT * px_per_pt))))
    header = image.crop((0, 0, width, header_h))
    footer = image.crop((0, height - footer_h, width, height))
    return header, footer


def _times_font(size: int) -> ImageFont.ImageFont:
    path = TIMES_BOLD if TIMES_BOLD.exists() else TIMES_REGULAR
    if path.exists():
        return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def _compose_header(fund_name: str, header: Image.Image) -> Image.Image:
    if fund_name == "Fund 1":
        return header
    meta = FUND_HEADERS.get(fund_name) or FUND_HEADERS["Fund 1"]
    lines = [str(meta["title"]), *list(meta["lines"])]
    out = header.copy()
    draw = ImageDraw.Draw(out)
    draw.rectangle(_ADDR_WHITEN, fill=(255, 255, 255))
    font = _times_font(_ADDR_FONT_PX)
    y = _ADDR_TOP_Y
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        text_w = bbox[2] - bbox[0]
        draw.text((_ADDR_RIGHT_X - text_w, y), line, font=font, fill=(28, 28, 28))
        y += _ADDR_LEADING
    return out


def _extract_logo_from_master() -> None:
    master, pdf_height = _load_master_page()
    header, _footer = _band_crops(master, pdf_height)
    width, height = header.size
    crop = header.crop((0, 0, int(width * 0.42), height))
    crop.save(LOGO_PATH, format="PNG")


def convert_to_a4(pdf_bytes: bytes) -> bytes:
    """Rebuild every page as A4, fitting source content into the body box."""
    return convert_to_page_setup(pdf_bytes, DEFAULT_PAGE_SETUP)


def convert_to_page_setup(pdf_bytes: bytes, setup_name: str) -> bytes:
    """Rebuild every page at the Word paper size, fitting content into the body."""
    if setup_name not in PAGE_SETUPS:
        raise ValueError(f"Unknown page setup: {setup_name}")
    from app.services.pdf_service import open_pdf

    width, height = PAGE_SETUPS[setup_name]
    body = content_box(width, height)
    source = open_pdf(pdf_bytes)
    output = fitz.open()
    try:
        for index in range(source.page_count):
            page = output.new_page(width=width, height=height)
            src_page = source[index]
            if src_page.get_contents():
                page.show_pdf_page(body, source, index, keep_proportion=True)
        return output.tobytes(deflate=True, garbage=4)
    finally:
        source.close()
        output.close()


def sniff_upload_kind(filename: str, data: bytes) -> str:
    """Classify an upload from magic bytes first, then the file extension."""
    suffix = Path(filename).suffix.lower()
    if data.startswith(b"%PDF"):
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        return "docx"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "doc"
    if suffix == ".pdf":
        return "pdf"
    if suffix in {".docx", ".docm"}:
        return "docx"
    if suffix == ".doc":
        return "doc"
    raise ValueError("Upload a PDF or Word document (.pdf, .docx, .doc).")


def convert_word_to_pdf(filename: str, data: bytes) -> bytes:
    """Turn a Word upload into PDF bytes. Prefer Word, then LibreOffice, then docx."""
    kind = sniff_upload_kind(filename, data)
    if kind == "pdf":
        return data
    suffix = ".docx" if kind == "docx" else ".doc"
    for converter in (_word_com_to_pdf, _libreoffice_to_pdf):
        pdf = converter(data, suffix)
        if pdf:
            return pdf
    if kind == "docx":
        return _docx_to_pdf(data)
    raise ValueError(
        "This .doc file could not be converted. Save it as .docx or PDF and upload again."
    )


def ingest_upload(filename: str, data: bytes) -> bytes:
    """Accept Word or PDF of any size and return A4 with the standard margins."""
    kind = sniff_upload_kind(filename, data)
    if kind in {"doc", "docx"}:
        data = convert_word_to_pdf(filename, data)
    elif kind != "pdf":
        raise ValueError("Upload a PDF or Word document (.pdf, .docx, .doc).")
    return convert_to_a4(data)


def _word_com_to_pdf(data: bytes, suffix: str) -> bytes | None:
    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore
    except ImportError:
        return None
    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp, f"source{suffix}").resolve()
            dest = Path(tmp, "out.pdf").resolve()
            src.write_bytes(data)
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            doc = word.Documents.Open(str(src), ReadOnly=True, AddToRecentFiles=False)
            try:
                doc.ExportAsFixedFormat(str(dest), 17)
            except Exception:
                doc.SaveAs(str(dest), FileFormat=17)
            doc.Close(False)
            doc = None
            if dest.exists() and dest.stat().st_size > 20:
                return dest.read_bytes()
    except Exception:
        return None
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass
    return None


def _libreoffice_to_pdf(data: bytes, suffix: str) -> bytes | None:
    soffice = shutil.which("soffice") or shutil.which("soffice.exe")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp, f"source{suffix}")
        src.write_bytes(data)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                check=True,
                capture_output=True,
                timeout=90,
            )
        except (OSError, subprocess.SubprocessError, subprocess.TimeoutExpired):
            return None
        dest = src.with_suffix(".pdf")
        if dest.exists() and dest.stat().st_size > 20:
            return dest.read_bytes()
    return None


def _docx_to_pdf(data: bytes) -> bytes:
    """Render .docx content onto full A4 pages. convert_to_a4 then fits the body."""
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise ValueError(
            "python-docx is required to read Word files. Install requirements.txt."
        ) from exc

    document = Document(io.BytesIO(data))
    blocks: list[tuple[str, bytes | str]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(("text", text))
        for run in paragraph.runs:
            blips = run._element.findall(
                ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
            )
            for blip in blips:
                rid = blip.get(qn("r:embed"))
                if not rid:
                    continue
                part = document.part.related_parts.get(rid)
                if part is not None and getattr(part, "blob", None):
                    blocks.append(("image", part.blob))
    for table in document.tables:
        for row in table.rows:
            line = "  |  ".join(cell.text.strip() for cell in row.cells).strip(" |")
            if line:
                blocks.append(("text", line))

    margin = 72.0
    box = fitz.Rect(margin, margin, A4_WIDTH - margin, A4_HEIGHT - margin)
    output = fitz.open()
    page = output.new_page(width=A4_WIDTH, height=A4_HEIGHT)
    y = box.y0 + 12
    line_h = 14
    if not blocks:
        return output.tobytes(deflate=True, garbage=4)
    for kind, payload in blocks:
        if kind == "image" and isinstance(payload, (bytes, bytearray)):
            try:
                image = Image.open(io.BytesIO(payload))
                iw, ih = image.size
            except OSError:
                continue
            width = min(box.width, 240)
            height = width * (ih / max(iw, 1))
            if y + height > box.y1:
                page = output.new_page(width=A4_WIDTH, height=A4_HEIGHT)
                y = box.y0 + 12
            page.insert_image(
                fitz.Rect(box.x0, y, box.x0 + width, y + height),
                stream=bytes(payload),
                keep_proportion=True,
            )
            y += height + 10
            continue
        for line in _wrap_line(str(payload), box.width, 11):
            if y > box.y1 - 12:
                page = output.new_page(width=A4_WIDTH, height=A4_HEIGHT)
                y = box.y0 + 12
            page.insert_text(
                (box.x0, y),
                line,
                fontsize=11,
                fontname="helv",
                color=(0.05, 0.05, 0.05),
            )
            y += line_h
        y += 8
    return output.tobytes(deflate=True, garbage=4)


def _wrap_line(text: str, width: float, fontsize: float) -> list[str]:
    max_chars = max(20, int(width / (fontsize * 0.5)))
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if len(trial) <= max_chars:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def apply_letterhead(pdf_bytes: bytes, fund_name: str) -> bytes:
    """Paint header + footer on every page using that fund's letterhead."""
    if fund_name not in FUNDS:
        raise ValueError(f"Unknown fund letterhead: {fund_name}")
    from app.services.pdf_service import open_pdf

    header_path, footer_path = render_letterhead_strips(fund_name)
    doc = open_pdf(pdf_bytes)
    try:
        for page in doc:
            body = content_box(page.rect.width, page.rect.height)
            page.insert_image(
                fitz.Rect(0, 0, page.rect.width, body.y0),
                filename=str(header_path),
                keep_proportion=False,
                overlay=True,
            )
            page.insert_image(
                fitz.Rect(0, body.y1, page.rect.width, page.rect.height),
                filename=str(footer_path),
                keep_proportion=False,
                overlay=True,
            )
        return doc.tobytes(deflate=True, garbage=4)
    finally:
        doc.close()


def render_letterhead_strips(fund_name: str) -> tuple[Path, Path]:
    """Crop header/footer from the fund's letterhead scan (Fund1.pdf as fallback)."""
    if fund_name not in FUNDS:
        raise ValueError(f"Unknown fund letterhead: {fund_name}")
    LETTERHEADS_DIR.mkdir(parents=True, exist_ok=True)
    source = letterhead_source_path(fund_name)
    if source.exists():
        header, footer = _crop_full_page_strips(source)
    else:
        master, pdf_height = _load_master_page()
        header, footer = _band_crops(master, pdf_height)
        header = _compose_header(fund_name, header)
    header_path = letterhead_header_path(fund_name)
    footer_path = letterhead_footer_path(fund_name)
    header.save(header_path, format="PNG")
    footer.save(footer_path, format="PNG")
    if not LOGO_PATH.exists():
        _extract_logo_from_master()
    return header_path, footer_path


def default_letterhead_boxes(pdf_width: float, pdf_height: float, image_width: int, image_height: int) -> tuple[dict, dict]:
    body = content_box(pdf_width, pdf_height)
    sx = image_width / pdf_width
    sy = image_height / pdf_height
    header = {"x": 0.0, "y": 0.0, "w": float(image_width), "h": body.y0 * sy}
    footer = {
        "x": 0.0,
        "y": body.y1 * sy,
        "w": float(image_width),
        "h": (pdf_height - body.y1) * sy,
    }
    return header, footer
